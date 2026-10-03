"""
仓库管理视图
"""
import logging
import io
from django.db import models
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from apps.core.exceptions import BusinessException
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    QuarantineCase, QuarantineApproval, QuarantineUnit,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    QuarantineCaseSerializer, QuarantineCaseCreateSerializer,
    QuarantineUnitSerializer, QuarantineReinspectionSerializer,
    QuarantineReinspectionCreateSerializer,
    QuarantineApprovalSerializer, QuarantineDecisionCreateSerializer,
    QuarantineRevokeSerializer, StockMovementSerializer,
    StockOutCreateSerializer, StockTransferCreateSerializer,
)
from . import services

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 货物 / 库存 ====================

def _parse_page(request):
    try:
        page = max(int(request.query_params.get('page', 1)), 1)
        page_size = max(int(request.query_params.get('page_size', 10)), 1)
    except (TypeError, ValueError):
        page, page_size = 1, 10
    return page, page_size


class GoodsListView(APIView):
    """货物列表视图（含隔离数量与可用数量）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        held_sq = models.Subquery(
            QuarantineUnit.objects.filter(
                case__goods=models.OuterRef('pk'), hold_active=True
            ).values('case__goods').annotate(
                total=models.Sum('hold_quantity')
            ).values('total'),
            output_field=models.DecimalField(max_digits=12, decimal_places=2)
        )
        queryset = Goods.objects.select_related(
            'variety__category__unit'
        ).annotate(_quarantined_total=held_sq).filter(is_active=True).order_by('-created_at')

        keyword = request.query_params.get('keyword')
        if keyword:
            queryset = queryset.filter(
                models.Q(name__icontains=keyword) | models.Q(code__icontains=keyword)
            )

        page, page_size = _parse_page(request)
        total = queryset.count()
        goods = queryset[(page - 1) * page_size:page * page_size]
        serializer = GoodsSerializer(goods, many=True)
        return success_response(data={
            'list': serializer.data, 'total': total,
            'page': page, 'page_size': page_size
        })


# ==================== 领用 / 转移（受有效隔离阻断） ====================

class StockOutListView(APIView):
    """领用记录列表 / 发起领用"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockOut.objects.select_related('goods', 'operator').all().order_by('-created_at')
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        page, page_size = _parse_page(request)
        total = queryset.count()
        records = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': StockOutSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        serializer = StockOutCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        try:
            stock_out = services.create_requisition(
                operator=request.user,
                goods=Goods.objects.get(pk=data['goods']),
                quantity=data['quantity'],
                receiver=data['receiver'],
                receiver_dept=data.get('receiver_dept', ''),
                remark=data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info(f"User {request.user.username} created stock-out {stock_out.id}")
        return success_response(data=StockOutSerializer(stock_out).data, message='领用申请已提交')


class StockOutCompleteView(APIView):
    """领用出库执行：出库瞬间再次校验有效隔离"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            stock_out = StockOut.objects.get(pk=pk)
        except StockOut.DoesNotExist:
            return error_response(message='领用记录不存在', code=404)
        try:
            stock_out = services.complete_requisition(
                stock_out=stock_out, operator=request.user
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=StockOutSerializer(stock_out).data, message='出库完成')


class StockTransferView(APIView):
    """库位转移：有效隔离数量必须被阻断在转移之外"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = StockTransferCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        try:
            movement = services.transfer_goods(
                operator=request.user,
                goods=Goods.objects.get(pk=data['goods']),
                quantity=data['quantity'],
                to_location=data['to_location'],
                remark=data.get('remark', ''),
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(data=StockMovementSerializer(movement).data, message='转移完成')


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related('goods', 'operator').all().order_by('-stock_in_time')
        page, page_size = _parse_page(request)
        total = queryset.count()
        records = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': StockInSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Warning.objects.select_related('goods').all().order_by('-created_at')
        page, page_size = _parse_page(request)
        total = queryset.count()
        records = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': WarningSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Approval.objects.select_related('stock_out', 'approver').all().order_by('-created_at')
        page, page_size = _parse_page(request)
        total = queryset.count()
        records = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': ApprovalSerializer(records, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })


# ==================== 隔离处置案件 ====================

class QuarantineCaseListView(APIView):
    """隔离处置案件列表 / 立案"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = QuarantineCase.objects.select_related('goods', 'discoverer').all().order_by('-created_at')
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)
        goods_id = request.query_params.get('goods')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        page, page_size = _parse_page(request)
        total = queryset.count()
        cases = queryset[(page - 1) * page_size:page * page_size]
        return success_response(data={
            'list': QuarantineCaseSerializer(cases, many=True).data,
            'total': total, 'page': page, 'page_size': page_size
        })

    def post(self, request):
        serializer = QuarantineCaseCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        try:
            case = services.open_case(
                operator=request.user,
                goods=Goods.objects.get(pk=data['goods']),
                affected_quantity=data['affected_quantity'],
                discovery_reason=data['discovery_reason'],
                reason_detail=data.get('reason_detail', ''),
                batch_no=data.get('batch_no', ''),
                stock_in=StockIn.objects.filter(pk=data.get('stock_in')).first() if data.get('stock_in') else None,
                quarantine_location=data.get('quarantine_location') or '隔离区',
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info(f"User {request.user.username} opened quarantine case {case.code}")
        return success_response(data=QuarantineCaseSerializer(case).data, message='隔离处置案件已建立')


class QuarantineCaseDetailView(APIView):
    """案件详情：发现原因、影响数量、复检、审批、移动与每单位处置状态一屏可追溯"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            case = QuarantineCase.objects.select_related(
                'goods', 'discoverer'
            ).get(pk=pk)
        except QuarantineCase.DoesNotExist:
            return error_response(message='隔离处置案件不存在', code=404)

        return success_response(data={
            'case': QuarantineCaseSerializer(case).data,
            'units': QuarantineUnitSerializer(
                case.units.all(), many=True
            ).data,
            'reinspections': QuarantineReinspectionSerializer(
                case.reinspections.select_related('inspector').all(), many=True
            ).data,
            'approvals': QuarantineApprovalSerializer(
                case.approvals.select_related('approver').all(), many=True
            ).data,
            'movements': StockMovementSerializer(
                case.movements.select_related('operator').all(), many=True
            ).data,
        })

    def delete(self, request, pk):
        """撤销案件（仅限尚无处置决定的案件）"""
        try:
            case = QuarantineCase.objects.get(pk=pk)
        except QuarantineCase.DoesNotExist:
            return error_response(message='隔离处置案件不存在', code=404)
        serializer = QuarantineRevokeSerializer(data=request.data)
        serializer.is_valid(raise_exception=False)
        try:
            services.revoke_case(
                case=case, operator=request.user,
                reason=serializer.validated_data.get('reason', '')
                if serializer.validated_data else ''
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(message='案件已撤销，已生成反向移动记录')


class QuarantineReinspectionView(APIView):
    """登记复检记录"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            case = QuarantineCase.objects.get(pk=pk)
        except QuarantineCase.DoesNotExist:
            return error_response(message='隔离处置案件不存在', code=404)
        serializer = QuarantineReinspectionCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        try:
            record = services.add_reinspection(
                case=case, inspector=request.user,
                result=data['result'],
                finding=data.get('finding', ''),
                qualified_quantity=data.get('qualified_quantity', 0) or 0,
                concession_quantity=data.get('concession_quantity', 0) or 0,
                unqualified_quantity=data.get('unqualified_quantity', 0) or 0,
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(
            data=QuarantineReinspectionSerializer(record).data,
            message='复检记录已登记'
        )


class QuarantineDecisionView(APIView):
    """提交审批意见（放行 / 让步接收 / 退回），触发最终库存移动"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            case = QuarantineCase.objects.get(pk=pk)
        except QuarantineCase.DoesNotExist:
            return error_response(message='隔离处置案件不存在', code=404)
        serializer = QuarantineDecisionCreateSerializer(data=request.data)
        if not serializer.is_valid():
            first_error = list(serializer.errors.values())[0][0]
            return error_response(message=str(first_error))
        data = serializer.validated_data
        reinspection = None
        if data.get('reinspection'):
            reinspection = case.reinspections.filter(pk=data['reinspection']).first()
            if reinspection is None:
                return error_response(message='复检记录不属于该案件')
        try:
            approval = services.decide(
                case=case, approver=request.user,
                decision=data['decision'], quantity=data['quantity'],
                target_location=data.get('target_location', ''),
                opinion=data.get('opinion', ''),
                reinspection=reinspection,
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        logger.info(
            f"User {request.user.username} decided {approval.decision} "
            f"on quarantine case {case.code}"
        )
        return success_response(
            data=QuarantineApprovalSerializer(approval).data,
            message='审批意见已落账，库存移动已生成'
        )


class QuarantineRevokeDecisionView(APIView):
    """撤销审批决定：生成反向动作，历史记录保留"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            approval = QuarantineApproval.objects.get(pk=pk)
        except QuarantineApproval.DoesNotExist:
            return error_response(message='审批决定不存在', code=404)
        serializer = QuarantineRevokeSerializer(data=request.data)
        serializer.is_valid(raise_exception=False)
        reason = serializer.validated_data.get('reason', '') if serializer.validated_data else ''
        try:
            services.revoke_decision(
                approval=approval, operator=request.user, reason=reason
            )
        except BusinessException as exc:
            return error_response(message=exc.message, code=exc.code)
        return success_response(message='决定已撤销，已生成反向库存移动')
